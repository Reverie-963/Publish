#!/usr/bin/env python3
"""Real HTTP and source-save tests in a disposable fixture, not UI screenshots."""
import sys, json, tempfile, threading, urllib.request, urllib.error, shutil, io, zipfile, os
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from serve_editor import make_server
from edit_documents import sha,word_model,word_package,patch_word,convert
from docx import Document
from docx.shared import Inches
from PIL import Image

root=Path(tempfile.mkdtemp(prefix='index-edit-test-'));(root/'.file-index').mkdir()
md=root/'notes.md';original='\ufeff# Heading\r\n\r\nText **bold** and $x^2$.\r\n\r\n[Link](asset.png)\r\n';md.write_bytes(original.encode())
Image.new('RGB',(80,40),'#a78bfa').save(root/'asset.png')
doc=Document();doc.add_heading('Editing test',0);p=doc.add_paragraph('Keep ');p.add_run('bold').bold=True;p.add_run(' intact.');doc.add_picture(str(root/'asset.png'),width=Inches(1));doc.add_heading('Table',1);doc.add_table(rows=2,cols=2).cell(0,0).text='Cell';doc.add_page_break();doc.add_heading('Second page',1);doc.add_paragraph('Unchanged paragraph.');doc.save(root/'word.docx')
empty=root/'empty.md';empty.write_text('')
(root/'legacy.doc').write_bytes(convert((root/'word.docx').read_bytes(),'.docx','.doc'))
records=[]
for i,p in enumerate([md,root/'word.docx',empty,root/'legacy.doc']):
 out=f'.file-index/previews/{i}/preview.html';(root/out).parent.mkdir(parents=True);(root/out).write_text('preview');records.append({'id':str(i),'path':p.name,'preview':out,'status':'preview'})
(root/'.file-index/manifest.json').write_text(json.dumps({'entry':'00 - 文件索引.html','files':records,'directories':[]}));(root/'00 - 文件索引.html').write_text('<html>index</html>')
server=make_server(root);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start();origin=f'http://127.0.0.1:{server.server_port}';base=origin+server.editor.prefix

def request(path,data=None,headers=None):
 h={'Origin':origin,'X-Edit-Key':server.editor.write_key,'Content-Type':'application/json'};h.update(headers or {})
 req=urllib.request.Request(base+path,data=json.dumps(data).encode() if data is not None else None,headers=h)
 try:
  with urllib.request.urlopen(req) as r:return r.status,r.read(),r.headers
 except urllib.error.HTTPError as e:return e.code,e.read(),e.headers
try:
 code,body,_=request('/.file-index/previews/0/preview.html');assert code==200 and b'data-edit-block' in body and b'edit-config' in body
 src,out,rev,model,_=server.editor.model('0');key=next(k for k,v in model['blocks'].items() if 'Text' in v['text']);b=model['blocks'][key]
 data={'id':'0','revision':rev,'block':key,'start':b['start'],'end':b['end'],'old':b['text'],'text':'Changed **bold** and $x^2$.\nNew line.'}
 code,body,_=request('/_edit',data);assert code==200,body;result=json.loads(body)
 expected=original.replace('Text **bold** and $x^2$.','Changed **bold** and $x^2$.\r\nNew line.')
 assert md.read_bytes()==expected.encode();assert (root/'.file-index/history/0'/f'{rev}.md').read_bytes()==original.encode()
 assert request('/_edit',data)[0]==409 # Stale revision cannot overwrite saved text.
 assert request('/_edit',data,{'Origin':'https://untrusted.invalid'})[0]==403
 assert request('/_edit',data,{'X-Edit-Key':'wrong'})[0]==403
 assert request('/.file-index/history/0/'+rev+'.md')[0]==403
 assert request('/../../etc/passwd')[0]==403
 # Word plain text edits preserve the package and each unaffected XML/media entry.
 src,out,rev,model,package=server.editor.model('1');key=next(k for k,v in model['blocks'].items() if v['text']=='Keep bold intact.');block=model['blocks'][key]
 before=(root/'word.docx').read_bytes();data={'id':'1','revision':rev,'block':key,'old':block['text'],'text':'Keep bold revised.\nNew line.'}
 code,body,_=request('/_edit',data);assert code==200,body
 after=(root/'word.docx').read_bytes();a=zipfile.ZipFile(io.BytesIO(before));b=zipfile.ZipFile(io.BytesIO(after));assert a.namelist()==b.namelist()
 for name in a.namelist():
  if name!='word/document.xml':assert a.read(name)==b.read(name),name
 edited=Document(io.BytesIO(after));assert any(r.bold and r.text=='bold' for p in edited.paragraphs for r in p.runs)
 assert edited.tables[0].cell(0,0).text=='Cell' and len(edited.inline_shapes)==1
 # Editing table-cell text uses the same source-preserving path.
 src,out,rev,model,package=server.editor.model('1');key=next(k for k,v in model['blocks'].items() if v['text']=='Cell')
 code,body,_=request('/_edit',{'id':'1','revision':rev,'block':key,'old':'Cell','text':'Updated cell'});assert code==200,body
 assert Document(root/'word.docx').tables[0].cell(0,0).text=='Updated cell'
 # An external update and a failed write retain the disk version.
 src,out,rev,model,_=server.editor.model('2');empty.write_text('External edit')
 assert request('/_edit',{'id':'2','revision':rev,'block':'block-0','old':'','start':0,'end':0,'text':'Overwrite'})[0]==409
 assert empty.read_text()=='External edit'
 # Empty documents can be written; adding paragraphs remains valid Markdown.
 empty.write_text('');src,out,rev,model,_=server.editor.model('2')
 assert request('/_edit',{'id':'2','revision':rev,'block':'block-0','old':'','start':0,'end':0,'text':'# New\n\nContent'})[0]==200
 # Source HTML is always sandboxed even when accessed directly.
 (root/'unsafe.html').write_text('<script>parent.bad=true</script>')
 assert request('/unsafe.html')[2]['Content-Security-Policy']=='sandbox'
 # Legacy DOC saves through a verified conversion with a binary original backup.
 src,out,rev,model,_=server.editor.model('3');key=next(k for k,v in model['blocks'].items() if v['text']=='Keep bold intact.')
 code,body,_=request('/_edit',{'id':'3','revision':rev,'block':key,'old':'Keep bold intact.','text':'Keep bold updated.'});assert code==200,body
 assert (root/'legacy.doc').read_bytes().startswith(bytes.fromhex('d0cf11e0a1b11ae1'))
 assert (root/'.file-index/history/3'/f'{rev}.doc').exists()
 # Persist a synthetic result for the required Word render/inspection step.
 review=Path('/tmp/index-edit-word-review');review.mkdir(exist_ok=True);shutil.copy2(root/'word.docx',review/'edited.docx')
 print(json.dumps({'passed':['Markdown BOM/CRLF/formulas/links','Word DOCX text/style/media/table preservation','legacy DOC save and backup','table text edits','source backups','revision conflicts and external changes','same-origin/capability enforcement','path traversal/history isolation','sandboxed source HTML','empty Markdown authoring'],'review_docx':str(review/'edited.docx')},indent=2))
finally:server.shutdown();server.server_close();shutil.rmtree(root)
