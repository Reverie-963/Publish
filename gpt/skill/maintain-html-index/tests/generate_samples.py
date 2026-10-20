"""Create synthetic file-preview fixtures only; never use a real course folder."""
from pathlib import Path
import argparse,tempfile
import csv,json,zipfile,wave,struct,math,subprocess
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PI
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Font
from openpyxl.chart import BarChart,Reference
from PIL import Image,ImageDraw
from reportlab.pdfgen import canvas
parser=argparse.ArgumentParser();parser.add_argument('output',type=Path);parser.add_argument('--legacy',action='store_true');args=parser.parse_args();r=args.output.resolve()
if r.exists() and any(r.iterdir()):parser.error('Use an empty fixture directory')
r.mkdir(parents=True,exist_ok=True)
for n in ['01 - 文档','02 - 表格','03 - 图片','04 - 音视频','05 - 文本与代码/嵌套目录 #1','06 - 其他']: (r/n).mkdir(parents=True,exist_ok=True)
im=Image.new('RGB',(960,540),'#f5f1ff');d=ImageDraw.Draw(im);d.rectangle((60,60,900,470),outline='#7c3aed',width=4);d.text((110,100),'FILE PREVIEW TEST',fill='#25232b',font_size=44);d.text((110,185),'Image / Office / PDF',fill='#7c3aed',font_size=30);d.rectangle((110,280,380,410),fill='#7c3aed');d.ellipse((520,260,700,440),fill='#c4b5fd')
for ext in ['png','jpg','webp','bmp','tiff']: im.save(r/'03 - 图片'/('图示 #1.'+ext))
im.save(r/'03 - 图片/动画.gif',save_all=True,append_images=[Image.new('RGB',(960,540),'#c4b5fd')],duration=500,loop=0)
(r/'03 - 图片/矢量.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="800" height="400"><rect width="800" height="400" fill="#f3efff"/><circle cx="220" cy="200" r="100" fill="#7c3aed"/><text x="380" y="210" font-size="36">SVG preview</text></svg>')
doc=Document();doc.add_heading('Document preview test',0);doc.add_paragraph('This sample checks headings, tables and embedded images.');doc.add_heading('Table',1);t=doc.add_table(rows=1,cols=2);t.rows[0].cells[0].text='Item';t.rows[0].cells[1].text='Value'
for a,b in [('Alpha','12'),('Beta','24')]:c=t.add_row().cells;c[0].text=a;c[1].text=b
doc.add_picture(str(r/'03 - 图片/图示 #1.png'),width=Inches(5));doc.add_page_break();doc.add_heading('Second page',1);doc.add_paragraph('The second page must be present in the preview.');doc.save(r/'01 - 文档/Word 示例.docx')
prs=Presentation()
for i in range(2):
 slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Slide preview '+str(i+1);slide.shapes.add_picture(str(r/'03 - 图片/图示 #1.png'),PI(1),PI(1.6),width=PI(8))
prs.save(r/'01 - 文档/演示文稿.pptx')
wb=Workbook();ws=wb.active;ws.title='数据';ws.append(['Item','Value']);ws.append(['Alpha',12]);ws.append(['Beta',24]);ws.append(['Total','=SUM(B2:B3)']);ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=18
for c in ws[1]:c.fill=PatternFill('solid',fgColor='7C3AED');c.font=Font(bold=True,color='FFFFFF')
chart=BarChart();chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=3),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=3));ws.add_chart(chart,'D2');s=wb.create_sheet('说明');s.merge_cells('A1:C1');s['A1']='Merged title';s.append(['Second sheet','Visible',0]);ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.print_area='A1:L18';wb.save(r/'02 - 表格/工作簿.xlsx')
for ext,sep in [('csv',','),('tsv','\t')]:
 with (r/'02 - 表格'/('数据.'+ext)).open('w',newline='') as f:csv.writer(f,delimiter=sep).writerows([['名称','数值'],['甲',12],['含逗号,和空格',24]])
c=canvas.Canvas(str(r/'01 - 文档/原版.pdf'))
for i in range(2):c.setFont('Helvetica',24);c.drawString(70,760,'PDF preview - page '+str(i+1));c.drawImage(str(r/'03 - 图片/图示 #1.png'),70,360,width=450,height=253);c.showPage()
c.save()
(r/'01 - 文档/说明.md').write_text('# Markdown 预览\n\n这里是完整正文，包含公式 $x^2+y^2=1$。\n\n## 同名标题\n\n[查看 Word](<Word 示例.docx>)\n\n![图示](<../03 - 图片/图示 %231.png>)\n\n## 同名标题\n\n|项目|数值|\n|---|---|\n|甲|12|\n\n<script>window.injected=true</script>\n')
texts={'示例.m':'% MATLAB 示例：数组、函数、字符串与转置\nfunction y = scale_signal(x, gain)\n    arguments\n        x double\n        gain (1,1) double = 2\n    end\n    label = "scaled signal";\n    samples = [1, 2, 3];\n    y = gain .* x;\n    column = samples\';\n    disp(label);\nend\n','说明.txt':'纯文本预览。\n第二行。','配置.json':json.dumps({'name':'示例','value':42},ensure_ascii=False,indent=2),'配置.yaml':'name: 示例\nvalue: 42\n','示例.py':'def add(a, b):\n    return a + b\n','示例.c':'int add(int a, int b) { return a + b; }\n'}
for n,v in texts.items():(r/'05 - 文本与代码/嵌套目录 #1'/n).write_text(v)
(r/'01 - 文档/已有网页.html').write_text('<!doctype html><meta charset="utf-8"><h1>已有 HTML 页面</h1><p>保留页面内容，在受限 iframe 中预览。</p><script>window.top.location="https://example.invalid"</script>')
with zipfile.ZipFile(r/'06 - 其他/压缩包.zip','w') as z:z.writestr('说明.txt','Archive preview');z.writestr('sub/data.csv','a,b\n1,2')
(r/'06 - 其他/未知格式.example').write_bytes(b'custom binary\x00\x01')
(r/'06 - 其他/损坏文件.docx').write_bytes(b'not a real office document')
with wave.open(str(r/'04 - 音视频/声音.wav'),'w') as w:
 w.setnchannels(1);w.setsampwidth(2);w.setframerate(22050);w.writeframes(b''.join(struct.pack('<h',int(2000*math.sin(2*math.pi*440*i/22050))) for i in range(22050)))
subprocess.run(['ffmpeg','-y','-loglevel','error','-i',str(r/'04 - 音视频/声音.wav'),str(r/'04 - 音视频/声音.mp3')],check=True)
subprocess.run(['ffmpeg','-y','-loglevel','error','-loop','1','-i',str(r/'03 - 图片/图示 #1.png'),'-t','2','-c:v','libx264','-pix_fmt','yuv420p',str(r/'04 - 音视频/视频.mp4')],check=True)
print(r)

if args.legacy:
 jobs=[('01 - 文档/Word 示例.docx','doc:MS Word 97'),('01 - 文档/Word 示例.docx','odt'),('01 - 文档/Word 示例.docx','rtf'),('01 - 文档/演示文稿.pptx','ppt:MS PowerPoint 97'),('01 - 文档/演示文稿.pptx','odp'),('02 - 表格/工作簿.xlsx','xls:MS Excel 97'),('02 - 表格/工作簿.xlsx','ods')]
 for file,format in jobs:
  with tempfile.TemporaryDirectory(prefix='index-legacy-') as d:
   src=r/file
   subprocess.run(['soffice','-env:UserInstallation='+Path(d).as_uri(),'--headless','--convert-to',format,'--outdir',str(src.parent),str(src)],check=True,capture_output=True,text=True,timeout=90)
