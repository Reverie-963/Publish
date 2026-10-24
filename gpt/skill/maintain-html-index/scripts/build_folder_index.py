#!/usr/bin/env python3
"""Portable offline file index; source files are never rewritten."""
import argparse, csv, hashlib, html, json, os, re, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from html_outline import html_preview
from native_opener import register_originals
from urllib.parse import quote, urlparse, unquote

MARKER='<!-- codex-folder-index v2 -->'
SKILL=Path(__file__).resolve().parent.parent
OFFICE={'.doc','.docx','.odt','.rtf','.ppt','.pptx','.odp','.xls','.ods'}
IMAGES={'.png','.jpg','.jpeg','.gif','.webp','.bmp','.tif','.tiff','.svg'}
AUDIO={'.mp3','.wav','.ogg','.m4a','.flac'}
VIDEO={'.mp4','.webm','.mov','.m4v','.avi'}
TEXT={'.txt','.json','.yaml','.yml','.xml','.py','.c','.h','.cpp','.m','.js','.ts','.css','.sh','.log','.tex','.ini','.toml'}
ESC=html.escape

def digest(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return h.hexdigest()

def rel(a,b): return '/'.join(quote(x,safe='') for x in Path(os.path.relpath(b,a.parent)).parts)
def natural(s): return [int(x) if x.isdigit() else x.casefold() for x in re.split(r'(\d+)',s)]
def write(p,text):
 p.parent.mkdir(parents=True,exist_ok=True)
 if p.exists() and p.read_text(encoding='utf8')==text:return
 temp=p.with_name(p.name+'.tmp');temp.write_text(text,encoding='utf8');temp.replace(p)
def command(args,timeout=120):
 r=subprocess.run([str(x) for x in args],capture_output=True,text=True,timeout=timeout)
 if r.returncode:raise RuntimeError((r.stderr or r.stdout or '转换失败')[-500:])
 return r

def page(title,body,out,root,index,original=""):
 css='\n'.join((SKILL/'assets'/name).read_text() for name in ['reader.css','preview.css','code.css','scrollbar.css','page-zoom.css','compare.css'])
 math=rel(out,root/'.file-index/assets/katex/katex.min.css')
 toc=''
 start=body.find('<button class="toc-toggle"')
 if start>=0:
  end=body.index('</aside>',start)+len('</aside>');toc,body=body[start:end],body[:start]+body[end:]
 match=re.match(r'\s*(<h1\b[^>]*>.*?</h1>)',body,re.S)
 heading=match.group(1) if match else '<h1>'+ESC(title)+'</h1>'
 if match:body=body[match.end():]
 header='<header class="reader-heading">'+heading+'<div class="reader-links"><a class="return-index" href="'+rel(out,index)+'">返回总索引</a>'+original+'</div></header>'
 return '<!doctype html>\n'+MARKER+'\n<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+ESC(title)+'</title><style>'+css+'</style><link rel="stylesheet" href="'+math+'"></head><body id="reader-top">'+toc+'<div class="layout"><article>'+header+body+'</article></div><script src="'+rel(out,root/'.file-index/assets/reader.js')+'"></script><script src="'+rel(out,root/'.file-index/assets/page-zoom.js')+'"></script><script src="'+rel(out,root/'.file-index/assets/reading-position.js')+'"></script><script src="'+rel(out,root/'.file-index/assets/file-catalog.js')+'"></script><script src="'+rel(out,root/'.file-index/assets/compare.js')+'"></script><script src="'+rel(out,root/'.file-index/assets/native-original.js')+'"></script></body></html>'

def notice(text):return '<p class="notice">'+ESC(text)+'</p>'
def table(rows):return '<div class="scroll"><table>'+''.join('<tr>'+''.join('<'+('th' if i==0 else 'td')+'>'+ESC(str(c) if c is not None else '')+'</'+('th' if i==0 else 'td')+'>' for c in row)+'</tr>' for i,row in enumerate(rows))+'</table></div>'

def outline(entries):
 # Entries are (target id, label, depth); omit navigation with no useful choice.
 if len(entries)<2:return ''
 links=''.join('<a href="#'+ESC(target,quote=True)+'" style="margin-left:'+str(min(depth,6)*12)+'px">'+ESC(label)+'</a>' for target,label,depth in entries)
 return '<button class="toc-toggle" aria-expanded="false" aria-controls="reader-toc">目录</button><aside class="toc-panel" id="reader-toc" inert><nav class="toc" aria-label="文稿目录">'+links+'</nav></aside>'

def pdf_outline(reader,limit,prefix='page'):
 entries=[]
 def visit(items,depth=0):
  for item in items:
   if isinstance(item,list):visit(item,depth+1);continue
   try:
    n=reader.get_destination_page_number(item)
    label=str(item.title).strip()
    if n is not None and 0<=n<limit and label:entries.append((prefix+'-'+str(n+1),label,depth))
   except (ValueError,TypeError,KeyError,AttributeError):continue
 try:visit(reader.outline)
 except (ValueError,TypeError,KeyError):pass
 # Keep page navigation as well: bookmarks may only cover some pages.
 pages=[(prefix+'-'+str(i+1),'第 '+str(i+1)+' 页',0) for i in range(limit)]
 return [(target,' / '.join(label for t,label,d in entries if t==target) or fallback,min([d for t,label,d in entries if t==target] or [0])) for target,fallback,depth in pages]

def image_dimensions(path):
 from PIL import Image
 with Image.open(path) as image:return ' width="'+str(image.width)+'" height="'+str(image.height)+'"'

def pdf_preview(pdf,out,max_pages,with_outline=True):
 from pypdf import PdfReader
 reader=PdfReader(str(pdf));count=len(reader.pages)
 if count==0:raise RuntimeError('PDF 没有页面')
 cmd=shutil.which('pdftoppm')
 if not cmd:raise RuntimeError('缺少 pdftoppm，无法制作保留版面的页面预览')
 for old in out.parent.glob('page-*.png'):old.unlink()
 command([cmd,'-f','1','-l',str(min(count,max_pages)),'-scale-to','1500','-png',pdf,out.parent/'page'],180)
 images=sorted(out.parent.glob('page-*.png'),key=lambda p:natural(p.name))
 if len(images)!=min(count,max_pages):raise RuntimeError('PDF 页面转换数量不完整')
 body=''.join('<figure class="page" id="page-'+str(i+1)+'"><img'+image_dimensions(p)+' loading="lazy" src="'+rel(out,p)+'" alt="第 '+str(i+1)+' 页"><figcaption>'+str(i+1)+' / '+str(count)+'</figcaption></figure>' for i,p in enumerate(images))
 notes=[]
 if count>max_pages:notes.append(f'共 {count} 页，仅预览前 {max_pages} 页；完整内容请打开原件。')
 return (outline(pdf_outline(reader,len(images))) if with_outline else '')+body,notes

def office_pdf(src,out):
 ext=src.suffix.lower()
 required={'.docx':'word/document.xml','.pptx':'ppt/presentation.xml','.xlsx':'xl/workbook.xml','.odt':'mimetype','.ods':'mimetype','.odp':'mimetype'}
 if ext in required:
  if not zipfile.is_zipfile(src):raise RuntimeError('文件扩展名与实际格式不符，无法预览')
  with zipfile.ZipFile(src) as z:
   if required[ext] not in z.namelist():raise RuntimeError('Office 文件结构不完整')
 if ext in {'.doc','.ppt','.xls'}:
  with src.open('rb') as f:
   if f.read(8)!=bytes.fromhex('d0cf11e0a1b11ae1'):raise RuntimeError('旧版 Office 文件头不匹配')
 binary=shutil.which('soffice') or shutil.which('libreoffice')
 if not binary:raise RuntimeError('缺少 LibreOffice；请使用可用的 Office 导出工具制作 PDF 后再预览')
 with tempfile.TemporaryDirectory(prefix='index-office-') as d:
  t=Path(d); command([binary,'-env:UserInstallation='+ (t/'profile').as_uri(),'--headless','--convert-to','pdf','--outdir',t,src],120)
  result=t/(src.stem+'.pdf')
  if not result.is_file():raise RuntimeError('Office 转换没有生成 PDF（文件可能损坏或不支持）')
  dest=out.parent/'document.pdf';shutil.copy2(result,dest);return dest

def workbook(src,out,max_cells):
 import openpyxl
 wb=openpyxl.load_workbook(src,data_only=False,read_only=False);cached=openpyxl.load_workbook(src,data_only=True,read_only=False)
 sections=[];notes=[];total=0;formulas=False
 for i,ws in enumerate(wb.worksheets):
  merges={};covered=set()
  for m in ws.merged_cells.ranges:
   merges[(m.min_row,m.min_col)]=(m.max_row-m.min_row+1,m.max_col-m.min_col+1)
   if (m.max_row-m.min_row+1)*(m.max_col-m.min_col+1)>max_cells:raise RuntimeError('合并区域超出预览上限')
   for row in range(m.min_row,m.max_row+1):
    for col in range(m.min_col,m.max_col+1):
     if (row,col)!=(m.min_row,m.min_col):covered.add((row,col))
  if total+ws.max_row*ws.max_column>max_cells:raise RuntimeError(f'表格超过 {max_cells} 个单元格预览上限，请打开原件或缩小范围')
  total+=ws.max_row*ws.max_column;rows=[]
  for row in ws.iter_rows():
   cells=[]
   for cell in row:
    if (cell.row,cell.column) in covered:continue
    v=cell.value;formula=''
    if cell.data_type=='f':
     formula=str(v);v=cached[ws.title][cell.coordinate].value
     if v is None:v=formula;formulas=True
    attrs=''
    if (cell.row,cell.column) in merges:
     rs,cs=merges[(cell.row,cell.column)];attrs=f' rowspan="{rs}" colspan="{cs}"'
    style=[]
    if cell.font.bold:style.append('font-weight:700')
    if cell.fill.patternType=='solid' and cell.fill.fgColor.type=='rgb':style.append('background:#'+cell.fill.fgColor.rgb[-6:])
    if cell.font.color and cell.font.color.type=='rgb':style.append('color:#'+cell.font.color.rgb[-6:])
    if cell.alignment.horizontal in ['left','center','right']:style.append('text-align:'+cell.alignment.horizontal)
    cells.append('<td'+attrs+' style="'+';'.join(style)+'"'+(' title="公式：'+ESC(formula,quote=True)+'"' if formula else '')+'>'+ESC('' if v is None else str(v))+'</td>')
   rows.append('<tr>'+''.join(cells)+'</tr>')
  if ws._charts or ws._images:notes.append(f'工作表「{ws.title}」含图表或图片，下方提供版面预览。')
  sections.append('<section class="sheet" id="sheet-'+str(i)+'"><h2>'+ESC(ws.title)+'</h2><div class="scroll"><table>'+''.join(rows)+'</table></div></section>')
 nav=outline([('sheet-'+str(i),ws.title,0) for i,ws in enumerate(wb.worksheets)]+([('print-layout','图表与打印版面',0)] if any('图表或图片' in n for n in notes) else []))
 wb.close();cached.close()
 if formulas:notes.append('部分公式没有已保存的计算结果，显示公式本身；表格视图不会重新计算。')
 return nav+''.join(sections),notes

def render(src,out,args,mapping):
 ext=src.suffix.lower();notes=[];status='preview';title='<h1>'+ESC(src.name)+'</h1>'
 if ext=='.md':
  r=subprocess.run(['node',str(SKILL/'scripts/render_markdown.cjs')],input=json.dumps({'source':str(src),'output':str(out),'mapping':mapping}),capture_output=True,text=True)
  if r.returncode:raise RuntimeError('Markdown 渲染失败：'+r.stderr[-350:])
  data=json.loads(r.stdout);body=data['body'];notes=data['issues'];title=''
 elif ext in OFFICE:
  pdf=office_pdf(src,out);body,notes=pdf_preview(pdf,out,args.max_pages)
 elif ext=='.pdf':body,notes=pdf_preview(src,out,args.max_pages)
 elif ext in {'.xlsx','.xlsm'}:
  body,notes=workbook(src,out,args.max_cells)
  if any('图表或图片' in s for s in notes):
   pdf=office_pdf(src,out);snap,extra=pdf_preview(pdf,out,args.max_pages,with_outline=False);body+='<details id="print-layout"><summary>查看图表与打印版面</summary>'+snap+'</details>';notes+=extra
 elif ext in {'.csv','.tsv'}:
  if src.stat().st_size>args.max_text_bytes:raise RuntimeError('文本表格超出预览大小上限')
  with src.open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f,delimiter='\t' if ext=='.tsv' else ','))
  if sum(map(len,rows))>args.max_cells:raise RuntimeError('文本表格超出单元格预览上限')
  body=table(rows)
 elif ext in IMAGES:
  target=src
  if ext in {'.tif','.tiff','.bmp'}:
   from PIL import Image,ImageSequence
   with Image.open(src) as im:
    pages=[]
    for i,fr in enumerate(ImageSequence.Iterator(im)):
     if i>=args.max_pages:notes.append('图片页数超出上限，仅预览前部分页面。');break
     p=out.parent/f'image-{i}.png';fr.convert('RGB').save(p);pages.append('<img id="image-'+str(i+1)+'"'+image_dimensions(p)+' loading="lazy" src="'+rel(out,p)+'" alt="第 '+str(i+1)+' 页">')
    body=outline([('image-'+str(i+1),'第 '+str(i+1)+' 页',0) for i in range(len(pages))])+'<div class="media-frame">'+''.join(pages)+'</div>'
  else:body='<div class="media-frame"><img src="'+rel(out,target)+'" alt="'+ESC(src.name,quote=True)+'"></div>'
 elif ext in AUDIO|VIDEO:
  tag='audio' if ext in AUDIO else 'video';body='<'+tag+' controls preload="metadata" src="'+rel(out,src)+'"></'+tag+'>'
  notes.append('播放能力取决于浏览器对该文件编码的支持；无法播放时请打开原件。')
 elif ext in {'.html','.htm'}:
  body=html_preview(src,out,args.max_text_bytes)
  notes.append('HTML 在受限预览中显示；脚本及提交操作未启用。')
 elif ext=='.zip':
  with zipfile.ZipFile(src) as z:rows=[['文件','大小（字节）']]+[[f.filename,f.file_size] for f in z.infolist()]
  body=table(rows);notes.append('仅展示压缩包目录，未解压或执行内部文件。')
 elif ext in TEXT:
  if src.stat().st_size>args.max_text_bytes:raise RuntimeError('文本超出预览大小上限，请打开原件')
  code=src.read_text(encoding='utf-8-sig')
  result=subprocess.run(['node',str(SKILL/'scripts/highlight_code.cjs')],input=json.dumps({'code':code,'filename':src.name}),capture_output=True,text=True)
  if result.returncode:raise RuntimeError('代码配色失败：'+result.stderr[-300:])
  body=result.stdout
 else:
  body=notice('此类型尚无内容预览，请使用“打开原件”。');status='unsupported';notes=['此类型尚无内容预览']
 return title+''.join(notice(x) for x in notes)+body,notes,status

def main():
 p=argparse.ArgumentParser();p.add_argument('root',type=Path);p.add_argument('--name',default='00 - 文件索引.html');p.add_argument('--title');p.add_argument('--exclude',action='append',default=[]);p.add_argument('--include-originals',action='store_true');p.add_argument('--force',action='store_true');p.add_argument('--links',type=Path);p.add_argument('--max-pages',type=int,default=200);p.add_argument('--max-cells',type=int,default=200000);p.add_argument('--max-text-bytes',type=int,default=2_000_000);args=p.parse_args()
 root=args.root.resolve();index=root/args.name
 if not root.is_dir():p.error('root must be an existing directory')
 if min(args.max_pages,args.max_cells,args.max_text_bytes)<1:p.error('preview limits must be positive')
 if Path(args.name).name!=args.name or not args.name.endswith('.html'):p.error('--name must be a local HTML filename')
 if index.exists() and MARKER not in index.read_text(encoding='utf8'):p.error('Refusing to overwrite an unrelated HTML entry')
 work=root/'.file-index';manifest=work/'manifest.json'
 if work.exists() and not manifest.exists():p.error('Existing .file-index has no ownership manifest; preserve it and choose a clean folder')
 old=json.loads(manifest.read_text()) if manifest.exists() else {};previous={x['path']:x for x in old.get('files',[]) if 'sha256' in x}
 if old and old.get('generator')!='codex-folder-index-v2':p.error('Unrecognized ownership manifest')
 work.mkdir(exist_ok=True);assets=work/'assets';assets.mkdir(exist_ok=True)
 # Claim only a new directory before generation, so an interrupted run is recoverable.
 if not manifest.exists():write(manifest,json.dumps({'generator':'codex-folder-index-v2','files':[]}))
 for name in ['reader.js','page-zoom.js','reading-position.js','compare.js','compare.css','native-original.js']:write(assets/name,(SKILL/'assets'/name).read_text())
 katex=SKILL/'node_modules/katex/dist'
 if katex.is_dir():
  (assets/'katex').mkdir(exist_ok=True);shutil.copy2(katex/'katex.min.css',assets/'katex/katex.min.css');shutil.copytree(katex/'fonts',assets/'katex/fonts',dirs_exist_ok=True)
  shutil.copy2(katex.parent/'LICENSE',assets/'katex/LICENSE')
 else:write(assets/'katex/katex.min.css','/* Markdown math dependencies unavailable. */')
 excluded={'node_modules','assets','source_assets','__pycache__',*args.exclude};sources=[];directories=[]
 for directory,dirs,files in os.walk(root):
  dirs[:]=sorted([d for d in dirs if not d.startswith('.') and d not in excluded and not (Path(directory)/d).is_symlink()],key=natural)
  if Path(directory)!=root:directories.append(Path(directory).relative_to(root).as_posix())
  for name in sorted(files,key=natural):
   f=Path(directory)/name
   if f.is_symlink() or name.startswith(('.', '~$')) or f==index or f.name==args.name+'.tmp':continue
   if not args.include_originals and name.endswith(('.original.md','.original.html')):continue
   if f.suffix=='.html' and f.with_suffix('.md').exists() and '<!-- codex-md-reader v1 -->' in f.read_text(errors='replace')[:250]:continue
   if name=='index.html' and ('codex-index-state:start' in f.read_text(errors='replace') or MARKER in f.read_text(errors='replace')):continue
   sources.append(f)
 native_links=register_originals(root,sources)
 mapping={str(f):str(work/'previews'/hashlib.sha256(str(f.relative_to(root)).encode()).hexdigest()[:16]/'preview.html') for f in sources}
 for directory in directories:mapping[str(root/directory)]=str(work/'folders'/hashlib.sha256(directory.encode()).hexdigest()[:16]/'index.html')
 version=hashlib.sha256((''.join(digest(SKILL/x) for x in ['scripts/build_folder_index.py','scripts/render_markdown.cjs','scripts/important_blocks.cjs','scripts/highlight_code.cjs','scripts/html_outline.py','package-lock.json','assets/reader.css','assets/preview.css','assets/code.css','assets/scrollbar.css','assets/page-zoom.css','assets/compare.css'])+str((args.name,args.max_pages,args.max_cells,args.max_text_bytes))).encode()).hexdigest()
 links_hash=hashlib.sha256(json.dumps(sorted([str(f.relative_to(root)) for f in sources]+directories)).encode()).hexdigest()
 records=[];reused=0
 for src in sources:
  sourcepath=src.relative_to(root).as_posix();out=Path(mapping[str(src)]);out.parent.mkdir(parents=True,exist_ok=True);sha=digest(src);oldrec=previous.get(sourcepath)
  if not args.force and old.get('version')==version and oldrec and oldrec['sha256']==sha and out.exists() and oldrec['status']!='error' and (src.suffix.lower()!='.md' or oldrec.get('links_hash')==links_hash):
   rec=oldrec;reused+=1
  else:
   try:body,notes,status=render(src,out,args,mapping)
   except Exception as e:notes=[str(e)];status='error';body='<h1>'+ESC(src.name)+'</h1>'+notice('未能生成内容预览：'+str(e))
   original='<a class="original-link" href="'+rel(out,src)+'"'+(' data-native-open="'+ESC(native_links[str(src)],quote=True)+'"' if str(src) in native_links else '')+' target="_blank" rel="noopener">打开原件</a>'
   write(out,page(src.name,body,out,root,index,original))
   rec={'id':out.parent.name,'path':sourcepath,'label':src.name,'type':src.suffix.lower().lstrip('.') or '无扩展名','source':rel(index,src),'preview':rel(index,out),'sha256':sha,'status':status,'notes':notes,'links_hash':links_hash if src.suffix.lower()=='.md' else None}
  preview_text=out.read_text()
  # Prefer the marked header link; legacy cached pages used a final footer anchor.
  original_match=re.search(r'<a class="original-link"[^>]*',preview_text)
  marker=original_match.end() if original_match else preview_text.rfind('>打开原件</a>')
  start=original_match.start() if original_match else preview_text.rfind('<a ',0,marker)
  if marker>=0 and start>=0:
   anchor=re.sub(r' data-native-open="[^"]*"','',preview_text[start:marker])
   if str(src) in native_links:anchor+=' data-native-open="'+ESC(native_links[str(src)],quote=True)+'"'
   preview_text=preview_text[:start]+anchor+preview_text[marker:]
  native_script='<script src="'+rel(out,assets/'native-original.js')+'"></script>'
  if native_script not in preview_text:preview_text=preview_text.replace('</body>',native_script+'</body>')
  write(out,preview_text)
  if digest(src)!=sha:raise RuntimeError('Source changed during build: '+str(src))
  records.append(rec)
 if args.links:
  for i,item in enumerate(json.loads(args.links.read_text())):
   if urlparse(item['url']).scheme not in {'http','https'}:raise ValueError('Online links must use http/https')
   out=work/'previews'/('link-'+str(i))/'preview.html'
   content='<h1>'+ESC(item['label'])+'</h1>'+notice('这是在线入口，没有生成离线正文。')+'<p><a href="'+ESC(item['url'],quote=True)+'" target="_blank" rel="noopener">打开在线页面</a></p>'
   write(out,page(item['label'],content,out,root,index))
   records.append({'id':'link-'+str(i),'path':item.get('folder','在线入口')+'/'+item['label'],'label':item['label'],'type':'链接','source':item['url'],'preview':rel(index,out),'status':'external','external':True,'notes':[]})
 # A script catalog works on file:// without fetch or cross-frame DOM access.
 catalog=[{'id':r['id'],'label':r['path'],'preview':r['preview']} for r in records if not r.get('external')]
 write(assets/'file-catalog.js','window.FILE_INDEX_CATALOG='+json.dumps(catalog,ensure_ascii=False).replace('<','\\u003c')+';')
 write(work/'compare.html',(SKILL/'assets/compare-template.html').read_text().replace('{{INDEX}}',ESC(rel(work/'compare.html',index),quote=True)))
 tree={}
 for directory in directories:
  branch=tree
  for folder in directory.split('/'):branch=branch.setdefault(folder,{})
 for f in records:
  branch=tree
  for folder in f['path'].split('/')[:-1]:branch=branch.setdefault(folder,{})
  branch.setdefault(None,[]).append(f)
 folders={}
 def collect(node,prefix=''):
  for k,v in node.items():
   if k is not None:
    key=prefix+k;folders[key]={'node':v,'page':work/'folders'/hashlib.sha256(key.encode()).hexdigest()[:16]/'index.html'};collect(v,key+'/')
 collect(tree)
 def branch_html(node,current,prefix=''):
  text=''
  for k,v in node.items():
   if k is None:
    for f in v:
     target=root/Path(unquote(f['preview']))
     text+='<li><a class="leaf" href="'+ESC(rel(current,target),quote=True)+'">'+ESC(f['label'])+'</a></li>'
   else:
    key=prefix+k;href=rel(current,folders[key]['page'])
    text+='<li><details><summary class="folder"><a class="folder-link" href="'+ESC(href,quote=True)+'" title="进入文件夹">'+ESC(k)+'</a></summary><ul>'+branch_html(v,current,key+'/')+'</ul></details></li>'
  return text
 template=(SKILL/'assets/navigation-template.html').read_text()
 def index_page(current,title,node,prefix='',nav=''):
  substitutions={'TITLE':ESC(title),'NAV_LINKS':nav,'TREE':branch_html(node,current,prefix),'SCROLLBAR_STYLE':(SKILL/'assets/scrollbar.css').read_text(),'PAGE_ZOOM_STYLE':(SKILL/'assets/page-zoom.css').read_text(),'PAGE_ZOOM_SCRIPT':rel(current,assets/'page-zoom.js'),'READING_POSITION_SCRIPT':rel(current,assets/'reading-position.js')}
  body=template.replace('<!doctype html>','<!doctype html>\n'+MARKER)
  body=body.replace('</head>','<link rel="stylesheet" href="'+rel(current,assets/'compare.css')+'"></head>').replace('</body>','<script src="'+rel(current,assets/'file-catalog.js')+'"></script><script src="'+rel(current,assets/'compare.js')+'"></script></body>')
  write(current,re.sub(r'\{\{(TITLE|NAV_LINKS|TREE|SCROLLBAR_STYLE|PAGE_ZOOM_STYLE|PAGE_ZOOM_SCRIPT|READING_POSITION_SCRIPT)\}\}',lambda m:substitutions[m.group(1)],body))
 index_page(index,args.title or root.name,tree)
 directory_records=[]
 for key,data in folders.items():
  current=data['page'];parent=key.rpartition('/')[0];up=folders[parent]['page'] if parent else index
  nav='<a href="'+rel(current,up)+'" data-parent>上一级</a><a href="'+rel(current,index)+'" data-root>总索引</a>'
  index_page(current,key.split('/')[-1],data['node'],key+'/',nav)
  directory_records.append({'path':key,'preview':rel(index,current)})
 # Maintain a previously generated alias only. Never replace an arbitrary index.html.
 legacy=root/'index.html'
 if legacy!=index and legacy.exists() and ('codex-index-state:start' in legacy.read_text(errors='replace') or MARKER in legacy.read_text(errors='replace')):
  write(legacy,'<!doctype html>'+MARKER+'<meta charset="utf-8"><meta http-equiv="refresh" content="0;url='+ESC(quote(args.name),quote=True)+'"><a href="'+quote(args.name)+'">打开文件索引</a>')
 options={key:getattr(args,key) for key in ['name','title','exclude','include_originals','max_pages','max_cells','max_text_bytes']}
 options['links']=os.path.relpath(args.links.resolve(),root) if args.links else None
 write(manifest,json.dumps({'generator':'codex-folder-index-v2','version':version,'entry':args.name,'build_options':options,'files':records,'directories':directory_records},ensure_ascii=False,indent=2))
 # Remove only retired generated folders explicitly owned by the previous manifest.
 keep={x['id'] for x in records}
 for record in old.get('files',[]):
  id=record.get('id','')
  if id not in keep and re.fullmatch(r'(?:[0-9a-f]{16}|link-[0-9]+)',id):
   retired=work/'previews'/id
   if retired.is_dir() and not retired.is_symlink():shutil.rmtree(retired)
 keep_dirs={x['preview'] for x in directory_records}
 for record in old.get('directories',[]):
  previous=record['preview']
  if previous not in keep_dirs:
   retired=root/unquote(previous);folder=retired.parent
   if folder.parent==work/'folders' and re.fullmatch(r'[0-9a-f]{16}',folder.name) and folder.is_dir() and not folder.is_symlink():shutil.rmtree(folder)
 print(json.dumps({'entry':str(index),'files':len(records),'reused':reused,'errors':[{'path':x['path'],'notes':x['notes']} for x in records if x['status']=='error'],'unsupported':[x['path'] for x in records if x['status']=='unsupported']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
